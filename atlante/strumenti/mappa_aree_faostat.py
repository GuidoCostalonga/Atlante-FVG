"""Costruisce la corrispondenza fra codici area FAOSTAT e codici ISO3 della Banca Mondiale.

Legge i dizionari originali FAOSTAT presenti nel repository (fogli *_AreaCodes: Area Code, M49 Code, Area)
e collega il codice M49 al codice ISO3 tramite Natural Earth (campi ISO_N3 e WB_A3/ISO_A3).
Uso: python3 mappa_aree_faostat.py <ne50.geojson> <uscita.json>
"""
import glob, json, os, sys
import openpyxl

QUI = os.path.dirname(os.path.abspath(__file__))
SORGENTE = os.path.join(os.path.dirname(os.path.dirname(QUI)), 'open data')

ne = json.load(open(sys.argv[1]))
m49_iso = {}
for f in ne['features']:
    p = f['properties']
    for n3 in (p.get('ISO_N3'), p.get('ISO_N3_EH'), p.get('UN_A3')):
        if n3 and n3 != '-99':
            for iso in (p.get('WB_A3'), p.get('ISO_A3_EH'), p.get('ISO_A3'), p.get('ADM0_A3')):
                if iso and iso != '-99':
                    m49_iso.setdefault(str(int(n3)), iso)
                    break
# correzioni note: Kosovo non ha M49, Taiwan usa 158
m49_iso.setdefault('158', 'TWN')

aree = {}
for f in sorted(glob.glob(os.path.join(SORGENTE, '*dizionari originali FAOSTAT.xlsx'))):
    wb = openpyxl.load_workbook(f, read_only=True)
    for s in wb.sheetnames:
        if not s.endswith('AreaCodes'):
            continue
        for r in wb[s].iter_rows(min_row=4, values_only=True):
            r = r[1:]
            if not r or not r[0]:
                continue
            cod, m49, nome = str(r[0]), str(r[1] or '').lstrip("'"), r[2]
            if '.' in m49 or not m49.isdigit():
                continue  # aggregati e aree non classificate restano senza ISO3
            iso = m49_iso.get(str(int(m49)))
            if iso:
                aree[cod] = iso
    wb.close()
json.dump(aree, open(sys.argv[2], 'w'), indent=0, sort_keys=True)
print(len(aree), 'codici area FAOSTAT collegati a un ISO3')
