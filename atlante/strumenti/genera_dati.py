"""Genera i dati dell'Atlante dai file del repository.

Legge i fogli DATI_NAZIONALI_WB e DIZ_INDICATORI degli Excel dei lotti 1 e 2
(Banca Mondiale, 150 indicatori WDI) contenuti negli archivi ZIP della cartella
"open data", unisce la geometria Natural Earth e scrive in atlante/dati:

  meta.json                 indicatori, territori, regioni, provenienza
  indicatori/<codice>.json  una serie annuale per territorio
  note/<codice>.json        note originali per osservazione, dove presenti
  mondo.topo.json           confini semplificati (generati a parte, vedi LEGGIMI)

Uso:
  python3 genera_dati.py                       estrae dagli ZIP (circa 10 minuti)
  python3 genera_dati.py --cache grezzi.json   riusa un'estrazione già fatta
  aggiungere --eurostat eurostat.json --eurostat-api eurostat_api.json --regio regio.json --insiemi <cartella estratti>
  per i dati NUTS (Eurostat e Commissione europea) e gli archivi parziali
"""
import argparse, glob, io, json, math, os, zipfile
import openpyxl
from etichette import E, TEMI, AGGREGATI_IT
import genera_extra

QUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(QUI))
SORGENTE = os.path.join(REPO, 'open data')
USCITA = os.path.join(REPO, 'atlante', 'dati')

REGIONI = [  # (chiave Natural Earth REGION_WB, nome, colore chiaro, colore scuro); ordine validato per daltonismo
    ('Sub-Saharan Africa', 'Africa subsahariana', '#eb6834', '#d95926'),
    ('Europe & Central Asia', 'Europa e Asia centrale', '#2a78d6', '#3987e5'),
    ('South Asia', 'Asia meridionale', '#eda100', '#c98500'),
    ('East Asia & Pacific', 'Asia orientale e Pacifico', '#e87ba4', '#d55181'),
    ('Latin America & Caribbean', 'America latina e Caraibi', '#008300', '#008300'),
    ('North America', 'America settentrionale', '#e34948', '#e66767'),
    ('Middle East & North Africa', 'Medio Oriente e Nord Africa', '#1baf7a', '#199e70'),
]


def estrai():
    archivi = sorted(glob.glob(os.path.join(SORGENTE, 'Lotto 1  *.zip')) +
                     glob.glob(os.path.join(SORGENTE, 'Lotto 2  *.zip')))
    obs, meta, nomi, tipo = {}, {}, {}, {}
    visti = set()
    for a in archivi:
        with zipfile.ZipFile(a) as z:
            for membro in z.namelist():
                base = os.path.basename(membro)
                if not base.startswith(('L01_', 'L02_')) or base in visti:
                    continue
                visti.add(base)
                wb = openpyxl.load_workbook(io.BytesIO(z.read(membro)), read_only=True)
                if 'DATI_NAZIONALI_WB' in wb.sheetnames:
                    agg = membro.startswith('AGGREGATI/')
                    for r in wb['DATI_NAZIONALI_WB'].iter_rows(min_row=4, values_only=True):
                        r = r[1:]
                        if not r or r[0] is None:
                            continue
                        iso, nome, cod, ind, anno, val, _, unita, stato, dec, nota, agg_f, url = r[:13]
                        nomi[iso] = nome
                        tipo[iso] = 'agg' if agg else 'paese'
                        m = meta.setdefault(cod, {'name': ind, 'unit': unita, 'upd': agg_f, 'url': url, 'rows': 0, 'vals': 0})
                        m['rows'] += 1
                        if val is not None:
                            m['vals'] += 1
                        obs.setdefault(cod, {}).setdefault(iso, {})[anno] = [val, stato, nota, dec]
                    for r in wb['DIZ_INDICATORI'].iter_rows(min_row=4, values_only=True):
                        r = r[1:]
                        if r[0] in meta and 'def' not in meta[r[0]]:
                            meta[r[0]].update({'def': r[3], 'org': r[4], 'unit2': r[2], 'dl': r[7]})
                wb.close()
    return {'obs': obs, 'meta': meta, 'names': nomi, 'kind': tipo}


def arrotonda(v, interi):
    if v is None:
        return None
    if interi:
        return int(v)
    if v == 0:
        return 0
    cifre = 6 - int(math.floor(math.log10(abs(v)))) - 1
    r = round(v, max(cifre, 0))
    return int(r) if r == int(r) and abs(r) < 1e15 else r


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--cache')
    p.add_argument('--geo', default=os.path.join(USCITA, 'mondo.topo.json'))
    p.add_argument('--eurostat', help='JSON prodotto da estrai_eurostat.py')
    p.add_argument('--eurostat-api', dest='eurostat_api', help='JSON prodotto da estrai_eurostat_api.py (indicatori regionali dall’API di Eurostat)')
    p.add_argument('--regio', help='JSON prodotto da estrai_regio.py (indici e serie regionali della Commissione europea, DG REGIO)')
    p.add_argument('--insiemi', help='cartella con i JSON prodotti da estrai_lotti.py (un file per insieme di dati)')
    a = p.parse_args()
    if a.cache and os.path.exists(a.cache):
        d = json.load(open(a.cache))
    else:
        d = estrai()
        if a.cache:
            json.dump(d, open(a.cache, 'w'))

    topo = json.load(open(a.geo))
    geo = {}
    for g in topo['objects'][next(iter(topo['objects']))]['geometries']:
        pr = g['properties']
        geo.setdefault(pr['id'], pr)
    nomi_reg = {r[0] for r in REGIONI}

    territori = []
    for iso, nome in sorted(d['names'].items()):
        k = d['kind'][iso]
        g = geo.get(iso)
        if k == 'agg':
            n_it = AGGREGATI_IT.get(iso, nome)
        else:
            n_it = (g or {}).get('n') or nome
        reg = (g or {}).get('r') if k == 'paese' else None
        territori.append({'id': iso, 'n': n_it, 'en': nome, 'k': k,
                          'r': reg if reg in nomi_reg else None, 'g': bool(g) and k == 'paese',
                          'c': (g or {}).get('c') if k == 'paese' else None})

    os.makedirs(os.path.join(USCITA, 'indicatori'), exist_ok=True)
    indicatori = []
    for cod in sorted(d['meta']):
        m = d['meta'][cod]
        serie = d['obs'][cod]
        anni = sorted({int(y) for t in serie.values() for y, v in t.items() if v[0] is not None})
        tutti = sorted({int(y) for t in serie.values() for y in t})
        y0, y1 = anni[0], anni[-1]
        interi = all(float(v[0]).is_integer() for t in serie.values() for v in t.values() if v[0] is not None)
        note = {}
        dati = {}
        for iso, t in serie.items():
            arr = [arrotonda(t.get(str(y), [None])[0], interi) for y in range(y0, y1 + 1)]
            if any(x is not None for x in arr):
                dati[iso] = arr
            for y, v in t.items():
                if v[1] or v[2]:
                    note.setdefault(iso, {})[y] = [v[1], v[2]]
        out = {'c': cod, 'y0': y0, 'y1': y1, 'd': dati}
        if note:
            os.makedirs(os.path.join(USCITA, 'note'), exist_ok=True)
            with open(os.path.join(USCITA, 'note', cod + '.json'), 'w') as f:
                json.dump(note, f, separators=(',', ':'), ensure_ascii=False)
        with open(os.path.join(USCITA, 'indicatori', cod + '.json'), 'w') as f:
            json.dump(out, f, separators=(',', ':'), ensure_ascii=False)
        tema, et, unita = E[cod]
        paesi_con_dati = sum(1 for iso in dati if d['kind'][iso] == 'paese')
        indicatori.append({
            'c': cod, 't': tema, 'n': et, 'u': unita, 'en': m['name'],
            'def': m.get('def'), 'org': m.get('org'), 'agg': m['upd'], 'url': m['url'], 'dl': m.get('dl'),
            'y0': y0, 'y1': y1, 'api0': tutti[0], 'api1': tutti[-1],
            'righe': m['rows'], 'valori': m['vals'], 'paesi': paesi_con_dati,
            'interi': interi, 'note': bool(note), 'lotto': 1 if cod in LOTTO1 else 2, 'fam': 'wb',
        })

    temi = dict(TEMI)
    archivi = []
    acquisizione = {'1': '2026-10-24', '2': '2026-10-24T21:49:29Z / 2026-10-24T21:50:09Z'}
    if a.eurostat:
        eu, nome_tema = genera_extra.eurostat(a.eurostat, USCITA, territori, arrotonda)
        indicatori += eu
        temi[genera_extra.TEMA_EU] = nome_tema
    if a.eurostat_api:
        eu_api, scaricato = genera_extra.eurostat_api(a.eurostat_api, USCITA, territori, arrotonda)
        indicatori += eu_api
        acquisizione['api'] = scaricato
        temi[genera_extra.TEMA_EU] = genera_extra.NOME_TEMA_EU_ESTESO
    if a.regio:
        regio, scaricato = genera_extra.eurostat_api(a.regio, USCITA, territori, arrotonda, lotto='regio')
        indicatori += regio
        acquisizione['regio'] = scaricato
    esclusi = []
    if a.insiemi:
        archivi, esclusi = genera_extra.insiemi(a.insiemi, USCITA, territori, arrotonda)

    meta = {
        'titolo': 'Atlante dei dati',
        'fonte': 'Banca Mondiale, World Development Indicators (WDI)',
        'acquisizione': acquisizione,
        'temi': temi,
        'archivi': archivi,
        'archivi_esclusi': esclusi,
        'regioni': [{'k': k, 'n': n, 'c': c, 'cs': cs} for k, n, c, cs in REGIONI],
        'indicatori': indicatori,
        'territori': territori,
    }
    with open(os.path.join(USCITA, 'meta.json'), 'w') as f:
        json.dump(meta, f, separators=(',', ':'), ensure_ascii=False)
    n = genera_extra.sintesi(USCITA, indicatori, territori)
    print(len(indicatori), 'indicatori;', len(territori), 'territori;', n, 'voci di panoramica')


# Indicatori del lotto 1 (gli altri 100 sono del lotto 2), dal foglio INDICATORI_WB dell'indice 1
LOTTO1 = set()

if __name__ == '__main__':
    import openpyxl as _o
    _wb = _o.load_workbook(os.path.join(SORGENTE, 'Open data internazionali  indice del lotto 1.xlsx'), read_only=True)
    for _r in _wb['INDICATORI_WB'].iter_rows(min_row=4, values_only=True):
        if _r[1] and _r[3] and _r[3] != 'Da acquisire':
            LOTTO1.add(_r[1])
    main()
