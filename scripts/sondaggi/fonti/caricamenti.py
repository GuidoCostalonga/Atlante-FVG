"""File caricati dal pannello di gestione (PDF, CSV, XLSX, immagini) e collegamenti aggiunti a mano.

Il pannello scrive nel ramo `sondaggi-dati`, cartella `caricati/`, un file JSON cifrato per ogni
caricamento: {nome, tipo, istituto, committente, titolo, inizio, fine, pubblicato, territorio, elezione,
url, contenuto (base64)}. Qui si legge il file, si estraggono tabelle e nota informativa e si propone una
rilevazione di livello «primaria» (il documento è dell'istituto o del committente). I valori estratti da un
PDF o da un'immagine entrano come «automatica» solo se superano i controlli; una tabella CSV/XLSX compilata
dall'utente entra come «verificata» perché è una trascrizione manuale dichiarata tale.
"""
from __future__ import annotations

import base64
import csv
import datetime as dt
import io
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from archivio import inserisci  # noqa: E402
from comune import pulisci  # noqa: E402
from estrai import controlli_voto, leggi_righe, nota_informativa, ocr_righe, righe_da_tabella, righe_ocr_in_valori, tabelle_pdf, testo_pdf  # noqa: E402


def _quando(c: dict) -> dt.date | None:
    for k in ('fine', 'pubblicato'):
        if c.get(k):
            try:
                return dt.date.fromisoformat(c[k])
            except ValueError:
                pass
    return None


def da_pdf(c: dict, dati: bytes, ocr: bool) -> tuple[list[dict], dict, str]:
    pagine = testo_pdf(dati)
    tutto = '\n'.join(pagine)
    nota = nota_informativa(tutto)
    domande = []
    quando = _quando(c)
    for tab in tabelle_pdf(dati):
        righe = leggi_righe(righe_da_tabella(tab['righe']), quando)
        valide = [r for r in righe if r['valore'] is not None]
        if len(valide) >= 3:
            riconosciute = sum(1 for r in valide if r['soggetto'])
            if riconosciute >= 3:
                domande.append({'tipo': 'voto', 'elezione': c.get('elezione') or 'politiche', 'turno': c.get('turno'), 'testo': 'Intenzioni di voto (tabella del documento)',
                                'unita': '%', 'denominatore': 'validi', 'risposte': valide, 'origine': 'tabella', 'controlli': controlli_voto(valide), 'rif': f'pagina {tab["pagina"]}'})
            else:
                domande.append({'tipo': 'altro', 'testo': c.get('titolo') or c.get('nome'), 'unita': '%', 'risposte': valide, 'origine': 'tabella', 'controlli': {}, 'rif': f'pagina {tab["pagina"]}'})
    if not any(d['tipo'] == 'voto' for d in domande):
        # tabelle «Nome valore» nel testo con disposizione conservata
        for n, t in enumerate(pagine, 1):
            if re.search(r'intenzioni di voto|orientamenti di voto|se oggi si votasse', t, re.I):
                righe = leggi_righe([re.sub(r'\s{2,}', ' ', r) for r in t.split('\n')], quando)
                valide = [r for r in righe if r['valore'] is not None and r['soggetto']]
                if len(valide) >= 4:
                    tutte = [r for r in righe if r['valore'] is not None]
                    domande.append({'tipo': 'voto', 'elezione': c.get('elezione') or 'politiche', 'turno': c.get('turno'), 'testo': 'Intenzioni di voto (testo del documento)',
                                    'unita': '%', 'denominatore': 'validi', 'risposte': tutte, 'origine': 'testo', 'controlli': controlli_voto(tutte), 'rif': f'pagina {n}'})
                    break
    if ocr and not any(d['tipo'] == 'voto' for d in domande):
        from fonti.swg import pagina_in_immagine
        for n, t in enumerate(pagine, 1):
            if re.search(r'intenzioni di voto|orientamenti di voto|se oggi si votasse', t, re.I) or (len(pagine) <= 3):
                png = pagina_in_immagine(dati, n)
                righe = ocr_righe(png) if png else []
                if righe:
                    risposte, _ = righe_ocr_in_valori(righe, quando)
                    voto = [r for r in risposte if r['valore'] is not None]
                    if sum(1 for r in voto if r['soggetto']) >= 4:
                        domande.append({'tipo': 'voto', 'elezione': c.get('elezione') or 'politiche', 'turno': c.get('turno'), 'testo': 'Intenzioni di voto (lettura automatica della pagina)',
                                        'unita': '%', 'denominatore': 'validi', 'risposte': voto, 'origine': 'ocr', 'controlli': controlli_voto(voto), 'rif': f'pagina {n}'})
                        break
    return domande, nota, tutto[:4000]


def da_tabella_testuale(c: dict, dati: bytes) -> list[dict]:
    """CSV o XLSX con colonne «soggetto, valore» (eventuali «domanda», «variazione», «min», «max»)."""
    righe = []
    nome = (c.get('nome') or '').lower()
    if nome.endswith('.xlsx'):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(dati), read_only=True, data_only=True)
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                righe.append([pulisci(str(x)) if x is not None else '' for x in r])
    else:
        testo = dati.decode('utf-8-sig', 'replace')
        dialetto = csv.Sniffer().sniff(testo[:2000], delimiters=';,\t') if testo.strip() else csv.excel
        righe = [[pulisci(x) for x in r] for r in csv.reader(io.StringIO(testo), dialetto)]
    if not righe:
        return []
    intest = [x.lower() for x in righe[0]]
    def col(*nomi):
        for n in nomi:
            if n in intest:
                return intest.index(n)
        return None
    cs, cv = col('soggetto', 'partito', 'risposta', 'nome'), col('valore', 'percentuale', '%')
    if cs is None or cv is None:
        return []
    cd, cvar, cmin, cmax = col('domanda'), col('variazione'), col('min', 'minimo'), col('max', 'massimo')
    quando = _quando(c)
    per_domanda: dict[str, list] = {}
    from soggetti import riconosci
    from comune import numero
    for r in righe[1:]:
        if len(r) <= max(cs, cv) or not r[cs]:
            continue
        testo_d = r[cd] if cd is not None and cd < len(r) and r[cd] else 'Intenzioni di voto'
        sid, motivo = riconosci(r[cs], quando)
        per_domanda.setdefault(testo_d, []).append({'testo': r[cs], 'soggetto': sid, 'motivo': motivo, 'valore': numero(r[cv]),
                                                    'variazione': numero(r[cvar]) if cvar is not None and cvar < len(r) else None,
                                                    'min': numero(r[cmin]) if cmin is not None and cmin < len(r) else None,
                                                    'max': numero(r[cmax]) if cmax is not None and cmax < len(r) else None})
    domande = []
    for testo_d, risposte in per_domanda.items():
        voto = 'voto' in testo_d.lower() or all(r['soggetto'] for r in risposte)
        domande.append({'tipo': 'voto' if voto else (c.get('tipoDomanda') or 'tema'), 'elezione': c.get('elezione') if voto else None, 'turno': c.get('turno'),
                        'testo': testo_d, 'unita': '%', 'denominatore': 'validi' if voto else 'totale', 'risposte': risposte, 'origine': 'manuale',
                        'controlli': controlli_voto(risposte) if voto else {'ok': True, 'motivi': []}})
    return domande


def caricamento_in_rilevazione(c: dict, ocr: bool) -> dict:
    dati = base64.b64decode(c.get('contenuto') or '') if c.get('contenuto') else b''
    nome = (c.get('nome') or '').lower()
    domande, nota, testo = [], {}, ''
    tipo_doc = 'pdf' if nome.endswith('.pdf') else ('csv' if nome.endswith(('.csv', '.xlsx', '.tsv')) else ('immagine' if nome.endswith(('.png', '.jpg', '.jpeg', '.webp')) else 'html'))
    manuale = False
    if dati and tipo_doc == 'pdf':
        domande, nota, testo = da_pdf(c, dati, ocr)
    elif dati and tipo_doc == 'csv':
        domande = da_tabella_testuale(c, dati)
        manuale = True
    elif dati and tipo_doc == 'immagine' and ocr:
        righe = ocr_righe(dati)
        testo = '\n'.join(r['testo'] for r in righe)[:4000]
        risposte, nota = righe_ocr_in_valori(righe, _quando(c))
        voto = [r for r in risposte if r['valore'] is not None]
        if sum(1 for r in voto if r['soggetto']) >= 4:
            domande.append({'tipo': 'voto', 'elezione': c.get('elezione') or 'politiche', 'turno': c.get('turno'), 'testo': 'Intenzioni di voto (lettura automatica dell\'immagine)',
                            'unita': '%', 'denominatore': 'validi', 'risposte': voto, 'origine': 'ocr', 'controlli': controlli_voto(voto)})
    ril = {
        'fonte': 'caricamenti', 'riferimento': c['id'], 'titolo': c.get('titolo') or c.get('nome') or 'Documento caricato',
        'istituto': c.get('istituto'), 'committente': c.get('committente'), 'editore': c.get('editore'),
        'inizio': c.get('inizio') or nota.get('inizio'), 'fine': c.get('fine') or nota.get('fine'), 'pubblicato': c.get('pubblicato'),
        'campione': c.get('campione') or nota.get('campione'), 'metodo': c.get('metodo') or nota.get('metodo'), 'erroreDichiarato': nota.get('erroreDichiarato'),
        'territorio': c.get('territorio') or {'tipo': 'Italia', 'nome': 'Italia', 'codice': 'IT'}, 'livello': c.get('livello') or 'primaria',
        'urls': [u for u in [c.get('url')] if u], 'documenti': [{'url': c.get('url') or f'caricati/{c["id"]}', 'tipo': tipo_doc, 'titolo': c.get('nome'), 'rif': 'file caricato dal pannello', 'testo': testo or None, 'caricato': c['id']}],
        'domande': domande, 'nota': c.get('nota'), 'caricatoDa': c.get('autore'), 'caricatoIl': c.get('quando'),
    }
    ril['criterioData'] = 'interviste' if ril['fine'] else 'pubblicazione'
    if c.get('elezione'):
        ril['elezione'] = c['elezione']
    if manuale and domande:
        ril['trascrizioneManuale'] = True
    return ril


def aggiorna(fonte: dict, stato: dict, archivio: dict, rete, opzioni: dict) -> dict:
    """I caricamenti vengono passati già decifrati da aggiorna.py in opzioni['caricamenti']."""
    esecuzione = opzioni['esecuzione']
    ocr = bool(opzioni.get('ocr', True))
    esito = {'letti': 0, 'nuovi': 0, 'aggiornati': 0, 'invariati': 0, 'coda': 0, 'protetti': 0, 'errori': [], 'finito': True}
    conteggio = {'nuova': 'nuovi', 'aggiornata': 'aggiornati', 'invariata': 'invariati', 'coda': 'coda', 'protetta': 'protetti'}
    fatti = set(stato.setdefault('avanzamento', {}).setdefault('fatti', []))
    for c in opzioni.get('caricamenti', []):
        if c['id'] in fatti:
            continue
        esito['letti'] += 1
        try:
            r = caricamento_in_rilevazione(c, ocr)
            es = inserisci(archivio, r, esecuzione)
            if r.get('trascrizioneManuale') and archivio['rilevazioni'].get(r['id'], {}).get('stato') in ('automatica', 'secondaria'):
                archivio['rilevazioni'][r['id']]['stato'] = 'verificata'
                archivio['rilevazioni'][r['id']]['protetta'] = True
            esito[conteggio[es]] += 1
            fatti.add(c['id'])
        except Exception as e:  # noqa: BLE001
            esito['errori'].append(f'{c.get("nome")}: {type(e).__name__}: {e}')
    stato['avanzamento']['fatti'] = sorted(fatti)
    return esito
