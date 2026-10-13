#!/usr/bin/env python3
"""Importa nell'archivio i fogli del file «Atlante FVG: archivio delle nuove raccolte» (raccolte del 12 e 13 ottobre 2026).

Il file Excel raccoglie fonti che l'archivio non aveva: registro regionale dei nidi, popolazione ISTAT per età,
consumo di suolo ISPRA, copertura internet AGCOM, appalti ANAC, medici di famiglia e pediatri, impianti fotovoltaici
e agrivoltaici del registro regionale, quotazioni immobiliari OMI, rendiconti dei comuni OpenBDAP 2018-2025, flussi
di cassa SIOPE, monitoraggio delle opere pubbliche BDAP-MOP, rendiconto 2025 di Drenchia e associazioni di
protezione civile, con i fogli di guida e di controllo della raccolta.

Per ogni foglio: legge titolo, descrizione e intestazione (riga 5; la colonna A è vuota), toglie le colonne di
contatto (telefoni, posta elettronica, PEC) e i codici fiscali dei fornitori privati, collega ogni riga al comune
con il codice ISTAT quando il foglio lo ha (altrimenti la pagina cerca il nome del comune nel testo), scrive i file
in dati/raccolte/ (JSON compresso con gzip e codificato in base64, come gli altri fogli, divisi in parti di 7 MB al
massimo) e aggiunge il foglio all'indice della pagina (const MAN) con ambito, nome leggibile, descrizione, fonte e
data di consultazione. Le colonne stanno nel primo file del foglio e i conteggi per comune sono un elenco nell'ordine
di MAN.istat, per non appesantire la pagina. I fogli già importati vengono sostituiti: lo script si può rieseguire.

Uso: python scripts/importa_raccolte.py <file.xlsx> [--cache CARTELLA] [--prova]
  --cache  cartella dove tenere i fogli già letti dal file Excel (la lettura dei 4,4 GB del file è lenta);
           senza, una cartella nella directory temporanea del sistema
  --prova  non scrive nulla: stampa che cosa farebbe
Richiede openpyxl (pip install openpyxl), solo per leggere il file Excel.
"""
import base64, datetime, gzip, json, math, os, re, sys, tempfile, zipfile
from collections import Counter
from multiprocessing import Pool
from pathlib import Path

RADICE = Path(__file__).resolve().parent.parent
PAGINA = RADICE / 'index.html'
CARTELLA = 'dati/raccolte'
MAX_FILE = 7 * 1024 * 1024
sys.path.insert(0, str(Path(__file__).resolve().parent))
from aggiorna_dati import riga_js, sostituisci_js, conta_testo  # noqa: E402

# colonne tolte per tutela dei privati (come nel resto dell'archivio)
PRIVATE = re.compile(r'telefono|e-?mail|\bpec\b|cellulare|\bfax\b', re.I)
PRIVATE_ESATTE = {'fornitore_cf'}
# colonne del codice ISTAT del comune, in ordine di preferenza
COL_CODICE = ['codice_istat', 'codice_istat_normalizzato', 'Codice comune', 'codice_istat_attuale', 'codici_istat_intersecati']
# fogli con lo stesso nome di fogli già nell'archivio
RINOMINA = {'INDICE': 'INDICE_raccolte_ottobre', 'LEGGIMI': 'LEGGIMI_raccolte_ottobre'}

URL = {
    'nidi': 'https://www.regione.fvg.it/rafvg/cms/RAFVG/famiglia-casa/politiche-famiglia/FOGLIA27/allegati/20261008_Strutture_per_registro_Regionale_Ottobre_26.xlsx',
    'suolo': 'https://www.isprambiente.gov.it/it/attivita/suolo-e-territorio/suolo/il-consumo-di-suolo/consumo_di_suolo_estratto_dati_2025_anni_2006_2024.xlsx',
    'agcom': 'https://geo.agcom.it/reportistica/desi/Survey_Italy_IT_DESI_2026_Data_2025_v2_pub.xlsx',
    'anac': 'https://data.open-contracting.org/en/publication/117',
    'medici': 'https://asfo.sanita.fvg.it/it/medici/',
    'aue': 'https://serviziogc.regione.fvg.it/geoserver/RIS_ENERG/wfs?version=1.0.0&service=wfs&request=GetFeature&typeName=RIS_ENERG:imp_ftv_agriftv&outputFormat=application/json&srsName=EPSG:4326',
    'omi': 'https://www1.agenziaentrate.gov.it/servizi/Consultazione/ricerca.htm',
    'istat': 'https://demo.istat.it/app/?i=POS',
    'bdap': 'https://openbdap.rgs.mef.gov.it/it/FET/analizza',
    'mop': 'https://bdap-opendata.rgs.mef.gov.it/opendata/spd_mop_prg_mon_reg06_01_9999',
    'siope_rgs': 'https://bdap-opendata.rgs.mef.gov.it/opendata/spd_rnd_spe_sio_reg06_01_{anno}',
    'siope_rgs_e': 'https://bdap-opendata.rgs.mef.gov.it/opendata/spd_rnd_ent_sio_reg06_01_{anno}',
    'siope_bi': 'https://www.siope.it/Siope/documenti/siope2/open/last/{file}',
    'drenchia': 'https://www.comune.drenchia.ud.it/media/files/030034/attachment/ALLEGATI_DELIBERA_CC_06.2026_APPROVAZIONE_CONSUNTIVO_2025.zip',
    'pc': 'https://www.protezionecivile.fvg.it/en/associazioni',
    'istat_confini': 'https://www.istat.it/storage/cartografia/confini_amministrativi/generalizzati/2026/Limiti01012026_g.zip',
}

GUIDA = 'Raccolte di ottobre 2026: guida e controlli'
A_NIDI = 'Nidi e prima infanzia'
A_ETA = 'Popolazione per età (ISTAT)'
A_SUOLO = 'Consumo di suolo (ISPRA)'
A_AGCOM = 'Copertura della rete internet (AGCOM)'
A_MEDICI = 'Medici di famiglia e pediatri'
A_APP = 'Appalti pubblici (ANAC)'
A_ENERGIA = 'Impianti fotovoltaici e agrivoltaici'
A_OMI = 'Valori immobiliari (OMI)'
A_SIOPE = 'Incassi e pagamenti dei comuni (SIOPE)'
A_MOP = 'Opere pubbliche: monitoraggio BDAP-MOP'
A_DRENCHIA = 'Rendiconto 2025 di Drenchia'
NUOVI_AMBITI = [GUIDA, A_NIDI, A_ETA, A_MEDICI, A_SUOLO, A_ENERGIA, A_AGCOM, A_OMI, A_APP, A_MOP, A_SIOPE] + \
    [f'Rendiconti dei comuni {a} (OpenBDAP)' for a in range(2025, 2017, -1)] + [A_DRENCHIA]

# id: (ambito, nome leggibile, descrizione, fonte, data di consultazione)
C12, C13 = '2026-10-12', '2026-10-13'
NOTA_RACCOLTA = 'Foglio della raccolta del 12 e 13 ottobre 2026.'
FOGLI = {
    'INDICE': (GUIDA, 'Indice delle raccolte di ottobre 2026', 'Elenco dei fogli della raccolta con il numero di righe e il contenuto di ciascuno. ' + NOTA_RACCOLTA, '', C13),
    'LEGGIMI': (GUIDA, 'Leggimi delle raccolte di ottobre 2026', 'Metodo, unità di misura, perimetro e avvertenze della raccolta: da leggere prima di usare i dati per confronti o pubblicazioni. ' + NOTA_RACCOLTA, '', C13),
    'PIANO_OPERATIVO': (GUIDA, 'Piano della raccolta', 'Le sette famiglie di fonti cercate, in ordine di priorità, con lo stato della raccolta (acquisito, parziale, non acquisito), il periodo, i fogli e i limiti.', '', C13),
    'COPERTURA_LIMITI': (GUIDA, 'Copertura e limiti', 'Che cosa è completo, che cosa è parziale e che cosa non è stato acquisito, con la fonte e il limite di ciascun tema.', '', C13),
    'FONTI': (GUIDA, 'Fonti delle raccolte', 'Indirizzi completi delle fonti, formato, periodo, stato dell\'acquisizione e condizioni di riutilizzo dichiarate o da verificare.', '', C13),
    'VERIFICA_QUALITA': (GUIDA, 'Controlli di qualità', 'Controlli di completezza e di collegamento ai comuni: valore atteso, valore osservato ed esito.', '', C13),
    'COMUNI_BASE': (GUIDA, 'I 215 comuni al 1° gennaio 2026', 'Codici ISTAT (Istituto nazionale di statistica), nomi e province dei 215 comuni al 1° gennaio 2026: la base con cui le raccolte collegano le righe ai comuni.', URL['istat_confini'], C12),
    'ALIAS_TERRITORIALI': (GUIDA, 'Nomi dei comuni nelle fonti', 'Regole con cui i nomi dei comuni scritti nelle fonti (forme brevi, nomi bilingui, comuni fusi) sono ricondotti al comune attuale. Il nome originale resta nei dati.', '', C12),
    'ANOMALIE': (GUIDA, 'Anomalie di localizzazione', 'Sedi non pubblicate e localizzazioni da verificare: le righe non sono state tolte. Da consultare prima di usare le sedi su una mappa.', '', C12),
    'DIZIONARIO': (GUIDA, 'Dizionario dei campi', 'Significato dei campi di ogni foglio della raccolta. Per ISPRA (Istituto superiore per la protezione e la ricerca ambientale) e AGCOM (Autorità per le garanzie nelle comunicazioni) ci sono anche le definizioni originali.', '', C13),
    'FONTI_DA_INTEGRARE': (GUIDA, 'Fonti che mancavano', 'Le fonti che mancavano all\'Atlante, che cosa è stato acquisito e che cosa resta da acquisire.', '', C13),
    'CATALOGO_INTEGRAZIONI': (GUIDA, 'Catalogo dei fogli integrati', 'Per ogni foglio integrato: righe importate, comuni presenti, periodo, file e prospetto di origine, limiti. Le righe sono un controllo dell\'importazione, non un totale economico.', '', C13),
    'CONTROLLI_INTEGRAZIONE': (GUIDA, 'Controlli delle integrazioni', 'Verifiche di copertura per comune e per periodo. La presenza di un comune non significa che ci siano tutti i mesi o tutti i prospetti.', '', C13),
    'RACCOLTA_13OTTOBRE': (GUIDA, 'Stato della raccolta al 13 ottobre 2026', 'Esito della prosecuzione del 13 ottobre 2026: rendiconto 2025 di Drenchia, SIOPE (Sistema informativo sulle operazioni degli enti pubblici) 2025 e 2026, rendiconto 2025 di Roveredo in Piano non acquisito, fonti ancora bloccate.', '', C13),
    'QA_PROSECUZIONE': (GUIDA, 'Controlli della prosecuzione', 'Controlli sui fogli del 13 ottobre: righe, comuni, mesi presenti, chiavi doppie e impronta dei file scaricati confrontata con quella pubblicata.', '', C13),
    'FONTI_PROSECUZIONE': (GUIDA, 'Fonti della prosecuzione', 'Fonti dei fogli del 13 ottobre 2026 con periodo e condizioni di riutilizzo. I dati SIOPE della Banca d\'Italia sono sotto licenza CC BY-NC 4.0 (riutilizzo non commerciale, citando la fonte).', '', C13),
    'APP_ESTRAZIONE': (A_APP, 'Appalti: controllo dell\'estrazione', 'Quante procedure nazionali sono state lette, quante selezionate per il Friuli Venezia Giulia e con quale criterio; prima e ultima data di inizio gara.', URL['anac'], C12),
    'NIDI_STRUTTURE': (A_NIDI, 'Nidi e servizi per la prima infanzia', 'Le 343 strutture del registro regionale dei servizi per la prima infanzia aggiornato all\'8 ottobre 2026: nidi d\'infanzia, spazi gioco, servizi domiciliari e centri per bambini e genitori, con tipologia, gestione, capacità ricettiva autorizzata, fascia d\'età, accreditamento e abbattimento rette. La capacità è quella autorizzata, non i posti liberi. Telefoni e indirizzi di posta elettronica sono stati tolti.', URL['nidi'], C12),
    'NIDI_COMUNI': (A_NIDI, 'Nidi per comune', 'Per ciascuno dei 215 comuni: servizi nel registro regionale, nidi d\'infanzia e somma della capacità autorizzata disponibile. Zero strutture significa nessuna struttura nel registro; una struttura non ha la capacità pubblicata ed è indicata a parte.', URL['nidi'], C12),
    'INFANZIA_DATI': (A_NIDI, 'Bambini di 0, 1 e 2 anni', 'Residenti di 0, 1 e 2 anni per comune al 1° gennaio di ogni anno dal 2019 al 2026, dalla popolazione ISTAT per età. Il 2026 è una stima.', URL['istat'], C13),
    'INFANZIA_0_2': (A_NIDI, 'Bambini da 0 a 2 anni per comune', 'Totale dei residenti da 0 a 2 anni per comune e anno, dal 2019 al 2026, sommando i tre anni di età ISTAT. Se manca un\'età la cella resta vuota. Il 2026 è una stima.', URL['istat'], C13),
    'SUOLO_COMUNI': (A_SUOLO, 'Consumo di suolo per comune', 'Serie ISPRA (Istituto superiore per la protezione e la ricerca ambientale) 2006-2024 per i 215 comuni: incremento netto, incremento lordo e ripristino per intervallo, suolo consumato nel 2024 in ettari e in percentuale (da 0 a 100). Pubblicazione 2025.', URL['suolo'], C12),
    'SUOLO_PROVINCE': (A_SUOLO, 'Consumo di suolo per provincia', 'Serie ISPRA 2006-2024 per le quattro province: incrementi per intervallo e suolo consumato nel 2024.', URL['suolo'], C12),
    'SUOLO_REGIONE': (A_SUOLO, 'Consumo di suolo in regione', 'Serie ISPRA 2006-2024 per il Friuli Venezia Giulia: incrementi per intervallo e suolo consumato nel 2024.', URL['suolo'], C12),
    'SUOLO_DIZIONARIO': (A_SUOLO, 'Consumo di suolo: definizioni', 'Definizioni e unità di misura dei campi, come pubblicate da ISPRA.', URL['suolo'], C12),
    'AGCOM_PROVINCE': (A_AGCOM, 'Copertura internet per provincia', 'Copertura delle reti fisse e mobili nelle quattro province al quarto trimestre 2025 (indicatori DESI, Indice di digitalizzazione dell\'economia e della società, di AGCOM, Autorità per le garanzie nelle comunicazioni): fibra fino a casa, rete fissa senza fili, DSL, 5G, reti ad altissima capacità, anche nelle zone rurali. Famiglie stimate dalla fonte; le quote sono frazioni da 0 a 1. Non sono dati comunali e non misurano abbonamenti o velocità. Il tabulato comunale di AGCOM non è stato acquisito: il collegamento ufficiale restituiva un errore.', URL['agcom'], C12),
    'AGCOM_DEFINIZIONI': (A_AGCOM, 'Copertura internet: definizioni', 'Definizioni delle tecnologie e della copertura dal file ufficiale AGCOM, nel testo originale in inglese.', URL['agcom'], C12),
    'MEDICI_SERVIZI': (A_MEDICI, 'Medici di famiglia e pediatri', 'Le 797 voci degli elenchi ufficiali delle tre aziende sanitarie (ASFO, Azienda sanitaria Friuli Occidentale; ASUFC, Azienda sanitaria universitaria Friuli Centrale; ASUGI, Azienda sanitaria universitaria Giuliano Isontina) al 12 ottobre 2026: medici di medicina generale e pediatri di libera scelta, con il distretto. Le 25 voci collettive (servizi, non singoli medici) sono indicate a parte. Non dicono se il medico accetta nuove scelte.', URL['medici'], C12),
    'AMBULATORI': (A_MEDICI, 'Ambulatori dei medici di famiglia e dei pediatri', 'Una riga per ogni sede indicata negli elenchi delle aziende sanitarie: indirizzo, comune, coordinate e controllo che le coordinate cadano nel comune dichiarato. Le sedi non pubblicate sono segnate come tali. Gli orari non sono stati raccolti; i numeri di telefono sono stati tolti.', URL['medici'], C12),
    'APP_ENTI': (A_APP, 'Appalti: enti appaltanti', 'Enti con sede in Friuli Venezia Giulia che hanno bandito procedure, dai dati ANAC (Autorità nazionale anticorruzione) ridistribuiti da Open Contracting Partnership. Selezione per sede dell\'ente, non per luogo dei lavori.', URL['anac'], C12),
    'APP_PROCEDURE': (A_APP, 'Appalti: procedure', 'Le 16.878 procedure di enti con sede in Friuli Venezia Giulia con inizio della gara tra il 1° febbraio 2018 e il 14 maggio 2024 (dati ANAC ridistribuiti da Open Contracting Partnership): oggetto, categoria, procedura, importo di gara, date, stato, numero di lotti, aggiudicazioni e contratti. Gli importi di gara non sono pagamenti. Archivio storico, non l\'elenco di tutti gli affidamenti in corso.', URL['anac'], C12),
    'APP_LOTTI': (A_APP, 'Appalti: lotti', 'I lotti delle procedure, collegati con il codice della procedura (OCID) e il CIG (Codice identificativo di gara). Gare, lotti, aggiudicazioni e contratti non si sommano fra loro.', URL['anac'], C12),
    'APP_AGGIUDICAZIONI': (A_APP, 'Appalti: aggiudicazioni', 'Le aggiudicazioni delle procedure con importo e data. L\'importo aggiudicato non è un pagamento.', URL['anac'], C12),
    'APP_FORNITORI': (A_APP, 'Appalti: fornitori aggiudicatari', 'Le imprese aggiudicatarie per ogni aggiudicazione, con la denominazione come scritta nella fonte (anche con caratteri guasti). I codici fiscali e le partite IVA delle imprese sono stati tolti.', URL['anac'], C12),
    'APP_CONTRATTI': (A_APP, 'Appalti: contratti', 'I contratti collegati alle aggiudicazioni, con importo e date. Non sono pagamenti.', URL['anac'], C12),
    'APP_TAPPE': (A_APP, 'Appalti: tappe e scadenze', 'Le tappe delle procedure registrate nella fonte (scadenze, date previste e realizzate).', URL['anac'], C12),
    'ENERGIA_AUE': (A_ENERGIA, 'Impianti fotovoltaici e agrivoltaici (registro AUE)', 'I 174 elementi del registro regionale degli impianti fotovoltaici e agrivoltaici a terra soggetti ad autorizzazione unica (AUE) al 12 ottobre 2026: tipo, progetto, potenza di picco e superficie come nella fonte (unità non dichiarate dal servizio), decreto, stato (in corso, autorizzato, costruito, in esercizio) e comuni interessati. Uno stesso progetto può avere più elementi. Non è il censimento di tutti gli impianti attivi del GSE (Gestore dei servizi energetici), che non è stato raggiungibile.', URL['aue'], C12),
    'ENERGIA_COMUNI': (A_ENERGIA, 'Impianti fotovoltaici e agrivoltaici per comune', 'Collegamento fra gli impianti del registro regionale e i comuni che intersecano. Un impianto può toccare più comuni: la potenza non va ripetuta nei totali comunali.', URL['aue'], C12),
    'ENERGIA_VERTICI': (A_ENERGIA, 'Impianti fotovoltaici e agrivoltaici: perimetri', 'I vertici originali dei perimetri degli impianti (longitudine e latitudine, sistema WGS84), in ordine, per ricostruirne la forma.', URL['aue'], C12),
    'OMI_COMUNI_FONTE': (A_OMI, 'Valori immobiliari: comuni della fonte', 'Le 225 unità comunali dell\'Osservatorio del mercato immobiliare (OMI) dell\'Agenzia delle entrate, compresi i comuni storici poi fusi, con il comune attuale.', URL['omi'], C12),
    'OMI_ZONE': (A_OMI, 'Valori immobiliari: zone', 'Le 639 zone OMI delle quattro province per il secondo semestre 2025, anche quelle senza quotazioni.', URL['omi'], C12),
    'OMI_QUOTAZIONI': (A_OMI, 'Valori immobiliari: quotazioni', 'Quotazioni OMI del secondo semestre 2025 per zona, destinazione (residenziale, commerciale, terziaria, produttiva), tipologia e stato di conservazione: compravendita in euro al metro quadro, affitto in euro al metro quadro al mese, superficie lorda (L) o netta (N). Sono intervalli indicativi, non stime di singoli immobili. Condizioni di riutilizzo dell\'Agenzia delle entrate da verificare.', URL['omi'], C12),
    'OMI_CONTROLLI': (A_OMI, 'Valori immobiliari: controllo delle interrogazioni', 'Registro delle interrogazioni per zona e destinazione. Zero righe significa nessuna quotazione pubblicata, non prezzo zero.', URL['omi'], C12),
    'PC_ASSOCIAZIONI': ('Associazioni', 'Associazioni di volontariato di protezione civile (raccolta 2026)', 'Le 76 associazioni dell\'elenco pubblicato dalla Protezione civile della Regione, con la specializzazione. Elenco letto dalla copia conservata della pagina ufficiale: la data dell\'elenco non è dichiarata e la verifica sul sito non è riuscita. Sede e comune non sono pubblicati e non sono stati ricavati dal nome.', URL['pc'], C13),
    'MOP_PROGETTI': (A_MOP, 'Opere pubbliche monitorate: progetti', 'Progetti di opere pubbliche in Friuli Venezia Giulia della banca dati di monitoraggio delle opere pubbliche (BDAP-MOP, Banca dati delle amministrazioni pubbliche, Ragioneria generale dello Stato): codice locale, CUP (Codice unico di progetto), titolare, natura, settore, date previste ed effettive. Le date 9999 non sono reali.', URL['mop'], C13),
    'MOP_PAGAMENTI': (A_MOP, 'Opere pubbliche monitorate: pagamenti', 'Pagamenti per progetto e anno, al netto dei recuperi. Non sono finanziamenti né costi realizzati.', URL['mop'], C13),
    'MOP_COSTI': (A_MOP, 'Opere pubbliche monitorate: piano dei costi', 'Piano dei costi per progetto e anno: importi da realizzare e realizzati. Non sono pagamenti.', URL['mop'], C13),
    'MOP_LOCALIZZAZIONI': (A_MOP, 'Opere pubbliche monitorate: localizzazioni', 'Comuni in cui si trovano i progetti. Un progetto può avere più localizzazioni: gli importi non si ripetono per ogni comune.', URL['mop'], C13),
    'SIOPE_BI_ENTI': (A_SIOPE, 'SIOPE: anagrafica dei comuni', 'I comuni del Friuli Venezia Giulia nell\'anagrafica SIOPE (Sistema informativo sulle operazioni degli enti pubblici) della Banca d\'Italia all\'8 ottobre 2026, compresi i codici dei comuni cessati. La data 9999-12-31 indica un ente ancora presente. Licenza CC BY-NC 4.0.', URL['siope_bi'].format(file='SIOPE_ANAGRAFICHE.zip'), C13),
    'SIOPE_BI_SPESE_2025': (A_SIOPE, 'SIOPE: pagamenti mensili dei comuni 2025', 'Pagamenti di ogni comune per mese e codice gestionale, da gennaio a dicembre 2025, dalla Banca d\'Italia (versione dell\'8 ottobre 2026). Importi del solo mese, non cumulati: la fonte li dà in centesimi di euro; l\'ultima colonna, aggiunta dall\'Atlante, li riporta in euro. Non vanno sommati con i fogli SIOPE della Ragioneria, che sono cumulati. Licenza CC BY-NC 4.0.', URL['siope_bi'].format(file='SIOPE_USCITE.2025.zip'), C13),
    'SIOPE_BI_SPESE_2026': (A_SIOPE, 'SIOPE: pagamenti mensili dei comuni 2026', 'Pagamenti di ogni comune per mese e codice gestionale, da gennaio a ottobre 2026 (ottobre parziale; anche i mesi precedenti possono essere rettificati), dalla Banca d\'Italia (versione dell\'8 ottobre 2026). Importi del solo mese, in centesimi di euro nella fonte; l\'ultima colonna, aggiunta dall\'Atlante, li riporta in euro. Licenza CC BY-NC 4.0.', URL['siope_bi'].format(file='SIOPE_USCITE.2026.zip'), C13),
    'SIOPE_BI_ENTRATE_2026': (A_SIOPE, 'SIOPE: incassi mensili dei comuni 2026', 'Incassi di ogni comune per mese e codice gestionale, da gennaio a ottobre 2026 (ottobre parziale), dalla Banca d\'Italia (versione dell\'8 ottobre 2026). Importi del solo mese, in centesimi di euro nella fonte; l\'ultima colonna, aggiunta dall\'Atlante, li riporta in euro. Licenza CC BY-NC 4.0.', URL['siope_bi'].format(file='SIOPE_ENTRATE.2026.zip'), C13),
    'SIOPE_BI_CODICI_SPESE': (A_SIOPE, 'SIOPE: codici dei pagamenti', 'Descrizioni dei codici gestionali dei pagamenti, con categoria e periodo di validità. Licenza CC BY-NC 4.0.', URL['siope_bi'].format(file='SIOPE_ANAGRAFICHE.zip'), C13),
    'SIOPE_BI_CODICI_ENTRATE': (A_SIOPE, 'SIOPE: codici degli incassi', 'Descrizioni dei codici gestionali degli incassi, con categoria e periodo di validità. Licenza CC BY-NC 4.0.', URL['siope_bi'].format(file='SIOPE_ANAGRAFICHE.zip'), C13),
    'DRENCHIA25_PAGINE': (A_DRENCHIA, 'Drenchia, rendiconto 2025: testo delle pagine', 'Il testo delle 173 pagine dei sei documenti del rendiconto 2025 del Comune di Drenchia (delibera del Consiglio comunale n. 6 del 2026), che manca nei dati regionali OpenBDAP. Estrazione automatica: per usare un valore controllare la pagina originale.', URL['drenchia'], C13),
    'DRENCHIA25_TABELLE': (A_DRENCHIA, 'Drenchia, rendiconto 2025: tabelle', 'Tutte le tabelle trovate nei sei documenti del rendiconto 2025 di Drenchia, con le celle originali. Estrazione automatica: per usare un valore controllare la pagina originale.', URL['drenchia'], C13),
    'DRENCHIA25_ENTRATE': (A_DRENCHIA, 'Drenchia, rendiconto 2025: entrate', 'Le entrate del rendiconto 2025 di Drenchia per titolo e tipologia, ricostruite dal documento ufficiale senza le righe di totale. Celle vuote non sostituite con zero. Estrazione automatica: per usare un valore controllare la pagina originale.', URL['drenchia'], C13),
    'DRENCHIA25_SPESE': (A_DRENCHIA, 'Drenchia, rendiconto 2025: spese', 'Le spese del rendiconto 2025 di Drenchia per missione, programma e titolo, ricostruite dal documento ufficiale senza le righe di totale. Celle vuote non sostituite con zero. Estrazione automatica: per usare un valore controllare la pagina originale.', URL['drenchia'], C13),
}
for anno in range(2019, 2027):
    FOGLI[f'ETA_{anno}'] = (A_ETA, f'Popolazione per età al 1° gennaio {anno}', f'Residenti di ogni comune al 1° gennaio {anno} per singolo anno di età (da 0 a 100 e oltre), sesso e stato civile, dalla rilevazione ISTAT (Istituto nazionale di statistica) della popolazione per età.' + (' Il 2026 è una stima e non ha lo stato civile.' if anno == 2026 else ''), f'{URL["istat"]}&a={anno}', C13)
for anno, file, mesi in ((2024, 'e', 'da gennaio a dicembre 2024'), (2025, 'e', 'da gennaio a dicembre 2025'), (2026, 'e', 'da gennaio a settembre 2026'),
                         (2024, 's', 'da gennaio a dicembre 2024'), (2025, 's', 'da gennaio a marzo 2025: la fonte non ha gli altri mesi'), (2026, 's', 'da gennaio a settembre 2026')):
    cosa = 'ENTRATE' if file == 'e' else 'SPESE'
    FOGLI[f'SIOPE_{cosa}_{anno}'] = (A_SIOPE, f'SIOPE: {"incassi" if file == "e" else "pagamenti"} cumulati dei comuni {anno} (Ragioneria)',
                                     f'{"Incassi" if file == "e" else "Pagamenti"} di ogni comune per mese e codice gestionale, {mesi}, dalla copia dei dati SIOPE (Sistema informativo sulle operazioni degli enti pubblici) pubblicata dalla Ragioneria generale dello Stato. L\'importo è cumulato da gennaio: i mesi non si sommano. Non va sommato con i fogli mensili della Banca d\'Italia.',
                                     URL['siope_rgs_e' if file == 'e' else 'siope_rgs'].format(anno=anno), C13)

SIGLE_BD = [(r'\bFPV\b', 'fondo pluriennale vincolato'), (r'\bFCDE\b', 'fondo crediti di dubbia esigibilità'), (r'\bUE\b', 'Unione europea')]


def nome_bd(foglio, origine):
    """«2024_Rendiconto SDB Allegato B FPV - Fondo Pluriennale Vincolato_FRIULI-VENEZIA GIULIA.csv» -> nome leggibile."""
    anno = '20' + foglio[2:4]
    p = origine.split(' / ')[-1]
    p = re.sub(r'_FRIULI-VENEZIA GIULIA\.csv$', '', p)
    p = re.sub(r'^\d{4}_', '', p)
    indicatori = ' PDI ' in f' {p} '
    p = re.sub(r'^(Bilanci di )?Rendiconto (SDB|PDI) ', '', p)
    for sigla, esteso in SIGLE_BD:
        if re.search(sigla, p):
            p = re.sub(r'\s*' + sigla + r'\s*-\s*', ' ', p) if re.search(sigla + r'\s*-\s*', p) else re.sub(sigla, esteso, p)
    p = re.sub(r'\s+-\s+', ': ', p).strip()
    p = p[0].upper() + p[1:]
    return f'Rendiconto {anno} · ' + ('Indicatori: ' if indicatori else '') + p


def descr_bd(foglio, origine, righe, comuni):
    anno = '20' + foglio[2:4]
    tipo = 'Piano degli indicatori del rendiconto' if foglio[5] == 'I' else 'Prospetto del rendiconto (schemi di bilancio)'
    return (f'{tipo} {anno} dei comuni del Friuli Venezia Giulia dalla Banca dati delle amministrazioni pubbliche (OpenBDAP) del Ministero '
            f'dell\'economia e delle finanze: {comuni} comuni. File di origine: {origine.split(" / ")[-1]}. Valori e codici come nella fonte; '
            f'competenza e cassa sono distinte; i prospetti diversi non si sommano fra loro. Un comune assente non significa zero.'
            + (' Nel 2025 mancano Drenchia (vedi i fogli del suo rendiconto) e Roveredo in Piano.' if anno == '2025' else ''))


# ---------------------------------------------------------------- lettura del file Excel
def _valore(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat(sep=' ')
    if isinstance(v, (datetime.date, datetime.time)):
        return v.isoformat()
    if isinstance(v, str):
        return v if v.strip() != '' else None
    if isinstance(v, float) and v.is_integer() and abs(v) < 2 ** 53:
        return int(v)
    return v


def _leggi_foglio(arg):
    xlsx, nome, cache = arg
    dest = Path(cache) / (nome + '.json.gz')
    if dest.exists():
        return nome
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[nome]
    ws.reset_dimensions()  # alcuni fogli dichiarano dimensioni sbagliate
    righe = [[_valore(v) for v in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    titolo = righe[1][1] if len(righe) > 1 and len(righe[1]) > 1 else None
    descr = righe[2][1] if len(righe) > 2 and len(righe[2]) > 1 else None
    h = next((i for i in range(3, len(righe)) if any(v is not None for v in righe[i])), None)
    cols, corpo = (list(righe[h][1:]), righe[h + 1:]) if h is not None else ([], [])
    while cols and cols[-1] is None:
        cols.pop()
    n = len(cols)
    rows = [r for r in ((list(r[1:]) + [None] * n)[:n] for r in corpo) if any(v is not None for v in r)]
    oltre = sum(1 for r in corpo if (r and r[0] is not None) or any(v is not None for v in r[1 + n:]))
    tmp = str(dest) + '.tmp'
    with gzip.open(tmp, 'wt', encoding='utf-8') as f:
        json.dump({'titolo': titolo, 'descr': descr, 'cols': cols, 'rows': rows, 'oltre': oltre, 'vuote_int': sum(c is None for c in cols),
                   'riga_intestazione': None if h is None else h + 1}, f, ensure_ascii=False, separators=(',', ':'))
    os.replace(tmp, dest)
    return nome


def fogli_del_file(xlsx):
    z = zipfile.ZipFile(xlsx)
    wbx = z.read('xl/workbook.xml').decode()
    rels = z.read('xl/_rels/workbook.xml.rels').decode()
    tgt = {}
    for rel in re.findall(r'<Relationship [^>]*>', rels):
        i, t = re.search(r'Id="([^"]+)"', rel).group(1), re.search(r'Target="([^"]+)"', rel).group(1)
        tgt[i] = t
    out = []
    for nome, rid in re.findall(r'<sheet [^>]*name="([^"]*)"[^>]*r:id="([^"]*)"', wbx):
        t = tgt[rid].lstrip('/')
        out.append((nome, z.getinfo(t if t.startswith('xl/') else 'xl/' + t).file_size))
    return out


def scrivi_dati(p, o):
    p = RADICE / p
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(base64.b64encode(gzip.compress(json.dumps(o, ensure_ascii=False, separators=(',', ':')).encode(), 9)).decode())


def codice6(v):
    s = str(v).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s.zfill(6) if s.isdigit() and len(s) <= 6 else None


def main():
    argomenti = sys.argv[1:]
    if not argomenti or argomenti[0].startswith('--'):
        raise SystemExit(__doc__)
    xlsx = argomenti[0]
    prova = '--prova' in argomenti
    cache = argomenti[argomenti.index('--cache') + 1] if '--cache' in argomenti else os.path.join(tempfile.gettempdir(), 'atlante_raccolte')
    Path(cache).mkdir(parents=True, exist_ok=True)
    elenco = fogli_del_file(xlsx)
    print(f'Fogli nel file: {len(elenco)}')
    with Pool(4) as p:
        for i, nome in enumerate(p.imap_unordered(_leggi_foglio, [(xlsx, n, cache) for n, _ in sorted(elenco, key=lambda x: -x[1])]), 1):
            if i % 50 == 0:
                print(f'  letti {i} fogli')
    leggi = lambda n: json.load(gzip.open(Path(cache) / (n + '.json.gz'), 'rt', encoding='utf-8'))
    catalogo = {r[0]: r for r in leggi('CATALOGO_INTEGRAZIONI')['rows']}

    html = PAGINA.read_text(encoding='utf-8')
    m_man, MAN = riga_js(html, 'MAN')
    istat = MAN['istat']
    pos = {c: i for i, c in enumerate(istat)}
    nomi_man = [None] * len(istat)
    _, DB = riga_js(html, 'DB')
    for c in DB['c']:
        if c['id'] in pos:
            nomi_man[pos[c['id']]] = c['n']
    ids_nuovi = {RINOMINA.get(n, n) for n, _ in elenco}
    vecchi = [f for f in MAN['manifest'] if f['id'] not in ids_nuovi]
    if len(vecchi) != len(MAN['manifest']):
        print(f'Fogli già importati che vengono sostituiti: {len(MAN["manifest"]) - len(vecchi)}')
    for f in MAN['manifest']:
        if f['id'] in ids_nuovi and not prova:
            for x in f['files']:
                (RADICE / x).unlink(missing_ok=True)
    nuovi, problemi, tolte = [], [], Counter()
    for nome, _ in elenco:
        d = leggi(nome)
        ident = RINOMINA.get(nome, nome)
        cols, rows = d['cols'], d['rows']
        if d['oltre'] or d['vuote_int']:
            problemi.append(f'{nome}: {d["oltre"]} righe con valori fuori dalle colonne, {d["vuote_int"]} intestazioni vuote')
        if not cols:
            problemi.append(f'{nome}: foglio senza intestazione, non importato')
            continue
        # colonne private
        tieni = [j for j, c in enumerate(cols) if not (c and (PRIVATE.search(str(c)) or str(c) in PRIVATE_ESATTE))]
        for j, c in enumerate(cols):
            if j not in tieni:
                tolte[f'{nome}.{c}'] += 1
        cols = [str(cols[j]) if cols[j] is not None else f'colonna {j + 1}' for j in tieni]
        rows = [[r[j] for j in tieni] for r in rows]
        # importi SIOPE della Banca d'Italia: colonna in euro accanto a quella originale in centesimi
        if nome.startswith('SIOPE_BI_') and 'importo_mese_centesimi_euro' in cols:
            jc = cols.index('importo_mese_centesimi_euro')
            cols.append('importo_mese_euro (calcolato: centesimi / 100)')
            for r in rows:
                v = r[jc]
                r.append(round(v / 100, 2) if isinstance(v, (int, float)) else None)
        # collegamento ai comuni
        jcod = next((cols.index(c) for c in COL_CODICE if c in cols), None)
        if jcod is not None:
            chiavi = []
            for r in rows:
                v = r[jcod]
                ks = [] if v is None else [k for k in (pos.get(codice6(x)) for x in re.split(r'[;,]', str(v))) if k is not None]
                ks = sorted(set(ks))
                chiavi.append(ks if len(ks) > 1 else (ks[0] if ks else -1))
            coll = f'colonna: {cols[jcod]}'
            if all(k == -1 for k in chiavi):  # colonna del codice vuota: si cerca il nome nel testo
                chiavi, coll, jcod = None, 'testo', None
        if jcod is not None:
            conteggio = Counter()
            for k in chiavi:
                for kk in (k if isinstance(k, list) else [k]):
                    if kk >= 0:
                        conteggio[str(kk)] += 1
        else:
            chiavi, coll = None, 'testo'
            conteggio = conta_testo(cols, rows, nomi_man)
        # conteggi per comune come elenco nell'ordine di MAN.istat: la pagina li legge con CONT[foglio][comune]
        conteggio = [int(conteggio.get(str(i), 0)) for i in range(len(istat))]
        # descrizione, nome, ambito, fonte, data
        if re.match(r'^BD\d\d_[SI]\d\d$', nome):
            c = catalogo.get(nome)
            origine = c[4] if c else ''
            amb, leggibile = f'Rendiconti dei comuni 20{nome[2:4]} (OpenBDAP)', nome_bd(nome, origine)
            descr, fonte, cons = descr_bd(nome, origine, len(rows), c[2] if c else '?'), URL['bdap'], C13
        elif nome in FOGLI:
            amb, leggibile, descr, fonte, cons = FOGLI[nome]
        else:
            problemi.append(f'{nome}: nessuna descrizione prevista, importato con quella del file')
            amb, leggibile, descr, fonte, cons = GUIDA, nome.replace('_', ' '), d['descr'] or '', '', C13
        # file divisi in parti di 7 MB al massimo
        prova_testo = len(base64.b64encode(gzip.compress(json.dumps({'rows': rows, 'k': chiavi}, ensure_ascii=False, separators=(',', ':')).encode(), 6)))
        parti = max(1, math.ceil(prova_testo / MAX_FILE * 1.05))
        passo = math.ceil(len(rows) / parti) if rows else 0
        files = [f'{CARTELLA}/{ident}_{j}.txt' for j in range(parti)]
        if not prova:
            for j, fn in enumerate(files):
                o = {'rows': rows[j * passo:(j + 1) * passo]}
                if j == 0:
                    o['cols'] = cols  # le colonne stanno nel file, non nell'indice della pagina, per non appesantirla
                if chiavi is not None:
                    o['k'] = chiavi[j * passo:(j + 1) * passo]
                scrivi_dati(fn, o)
        voce = {'id': ident, 'ambito': amb, 'righe': len(rows), 'ncol': len(cols), 'descr': descr, 'fonte': fonte, 'cons': cons, 'files': files, 'coll': coll, 'nome': leggibile}
        nuovi.append(voce)
        MAN['conteggi'][ident] = conteggio
    for f in MAN['manifest']:
        if f['id'] in ids_nuovi and f['id'] not in {v['id'] for v in nuovi}:
            MAN['conteggi'].pop(f['id'], None)
    MAN['manifest'] = vecchi + nuovi
    print(f'Fogli importati: {len(nuovi)}, righe: {sum(v["righe"] for v in nuovi):,}'.replace(',', '.'))
    print(f'Collegati con il codice ISTAT: {sum(1 for v in nuovi if v["coll"] != "testo")}; con il nome nel testo: {sum(1 for v in nuovi if v["coll"] == "testo")}')
    print('Colonne tolte per tutela dei privati:', ', '.join(sorted(tolte)) or 'nessuna')
    for p in problemi:
        print('ATTENZIONE', p)
    if prova:
        return 0
    html = sostituisci_js(html, m_man, MAN)
    # ordine degli ambiti nell'elenco dei fogli
    m = re.search(r"^const ORDINE_AMBITI = (\[.*\]);$", html, re.M)
    ordine = json.loads(m.group(1).replace("'", '"'))
    ordine = [a for a in ordine if a not in NUOVI_AMBITI]
    i = ordine.index('Regione in cifre 2026: integrazioni') if 'Regione in cifre 2026: integrazioni' in ordine else len(ordine)
    ordine = ordine[:i] + NUOVI_AMBITI + ordine[i:]
    html = html[:m.start(1)] + '[' + ', '.join("'" + a.replace("'", "\\'") + "'" for a in ordine) + ']' + html[m.end(1):]
    PAGINA.write_text(html, encoding='utf-8')
    print('Indice della pagina aggiornato.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
